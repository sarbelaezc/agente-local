from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class WorkbookError(ValueError):
    pass


@dataclass(frozen=True)
class ReviewedInvoice:
    company: str
    legal_name: str
    issued: date
    amount: Decimal
    number: str
    currency: str
    legacy_row: int | None = None
    confirmed_new: bool = False

    def validate(self) -> None:
        for value in (self.company, self.legal_name, self.number):
            if not value.strip() or len(value) > 500 or any(ord(c) < 32 for c in value):
                raise WorkbookError("Company, legal name and invoice number must be valid text")
        if self.currency != "COP":
            raise WorkbookError("Only confirmed COP amounts can be stored in the existing Valor column")
        if not self.amount.is_finite() or self.amount <= 0:
            raise WorkbookError("Invoice amount must be positive and finite")
        if self.amount != self.amount.quantize(Decimal("0.01")):
            raise WorkbookError("Invoice amount has more than two decimal places")
        if not isinstance(self.issued, date):
            raise WorkbookError("An invoice issue date is required")


@dataclass(frozen=True)
class RecordResult:
    number: str
    row: int
    inserted: bool


def _normalized(text: object) -> str:
    return " ".join(str(text or "").split()).casefold()


def _cell_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _matches_row(sheet: Worksheet, row: int, invoice: ReviewedInvoice) -> bool:
    amount = sheet.cell(row, 5).value
    return (
        _normalized(sheet.cell(row, 3).value) == _normalized(invoice.legal_name)
        and _cell_date(sheet.cell(row, 4).value) == invoice.issued
        and amount is not None
        and Decimal(str(amount)) == invoice.amount
    )


class InvoiceWorkbook:
    def __init__(self, path: Path, backup_dir: Path, sheet_name: str = "Hoja1") -> None:
        self.path = path
        self.backup_dir = backup_dir
        self.sheet_name = sheet_name

    def _sheet(self, workbook: Workbook) -> Worksheet:
        if self.sheet_name not in workbook.sheetnames:
            raise WorkbookError("Configured invoice sheet is missing")
        sheet = workbook[self.sheet_name]
        if not isinstance(sheet, Worksheet):
            raise WorkbookError("Configured invoice sheet is not a worksheet")
        headers = [sheet.cell(2, column).value for column in range(2, 6)]
        if headers != ["Empresa", "Raz\u00f3n Social", "Fecha", "Valor"]:
            raise WorkbookError("Workbook headers differ from the reviewed B2:E2 schema")
        number_header = sheet.cell(2, 6).value
        if number_header not in (None, "N\u00famero de factura"):
            raise WorkbookError("Column F is occupied; refusing to overwrite it")
        if number_header is None and any(
            sheet.cell(row, 6).value is not None for row in range(3, sheet.max_row + 1)
        ):
            raise WorkbookError("Column F contains data without the expected header")
        return sheet

    def legacy_matches(self, invoice: ReviewedInvoice) -> tuple[int, ...]:
        workbook = load_workbook(self.path)
        try:
            sheet = self._sheet(workbook)
            return tuple(
                row for row in range(3, sheet.max_row + 1)
                if sheet.cell(row, 6).value is None and _matches_row(sheet, row, invoice)
            )
        finally:
            workbook.close()

    def company_suggestions(self) -> dict[str, str]:
        """Maps normalized legal names to the company used in every historical row.

        Legal names recorded under more than one company are left out, so the
        reviewer decides instead of the agent guessing.
        """
        workbook = load_workbook(self.path)
        try:
            sheet = self._sheet(workbook)
            seen: dict[str, set[str]] = {}
            for row in range(3, sheet.max_row + 1):
                legal_name = _normalized(sheet.cell(row, 3).value)
                company = " ".join(str(sheet.cell(row, 2).value or "").split())
                if legal_name and company:
                    seen.setdefault(legal_name, set()).add(company)
            return {name: next(iter(companies)) for name, companies in seen.items() if len(companies) == 1}
        finally:
            workbook.close()

    def registration_status(
        self, number: str, legal_name: str, issued: date | None,
        amount: Decimal | None, currency: str,
    ) -> str:
        """Read-only evidence for classification; historical matches do not authorize moves."""
        if not number.strip() or not legal_name.strip() or issued is None or amount is None:
            return "incomplete"
        if currency != "COP" or not amount.is_finite() or amount <= 0:
            return "incomplete"
        workbook = load_workbook(self.path)
        try:
            sheet = self._sheet(workbook)
            matches = [
                row for row in range(3, sheet.max_row + 1)
                if _normalized(sheet.cell(row, 3).value) == _normalized(legal_name)
                and _normalized(sheet.cell(row, 6).value) == _normalized(number)
            ]
            if len(matches) > 1:
                return "conflict"
            invoice = ReviewedInvoice(legal_name, legal_name, issued, amount, number, currency)
            if matches:
                return "registered" if _matches_row(sheet, matches[0], invoice) else "conflict"
            historical = any(
                sheet.cell(row, 6).value is None and _matches_row(sheet, row, invoice)
                for row in range(3, sheet.max_row + 1)
            )
            return "historical_review" if historical else "not_registered"
        finally:
            workbook.close()

    def record(self, invoices: list[ReviewedInvoice]) -> list[RecordResult]:
        """Commit only reviewed records, with backup, conflict detection and verification."""
        if not invoices:
            raise WorkbookError("No approved invoices were supplied")
        for invoice in invoices:
            invoice.validate()
        if self.path.suffix.lower() != ".xlsx":
            raise WorkbookError("Only the reviewed XLSX workbook format is supported")
        if self.path.with_name("~$" + self.path.name).exists():
            raise WorkbookError("Close the workbook in Excel before approving a batch")
        original_hash = hashlib.sha256(self.path.read_bytes()).hexdigest()
        workbook = load_workbook(self.path)
        temporary: Path | None = None
        try:
            sheet = self._sheet(workbook)
            original_values = {
                (name, cell.coordinate): cell.value
                for name in workbook.sheetnames
                if isinstance(workbook[name], Worksheet)
                for row in workbook[name].iter_rows()
                for cell in row if cell.value is not None
            }
            sheet.cell(2, 6, "N\u00famero de factura")
            sheet.cell(2, 6)._style = copy(sheet.cell(2, 5)._style)
            results = []
            for invoice in invoices:
                exact = [
                    row for row in range(3, sheet.max_row + 1)
                    if _normalized(sheet.cell(row, 3).value) == _normalized(invoice.legal_name)
                    and _normalized(sheet.cell(row, 6).value) == _normalized(invoice.number)
                ]
                if len(exact) > 1:
                    raise WorkbookError("Existing invoice number is duplicated in the workbook")
                if exact:
                    row = exact[0]
                    if not _matches_row(sheet, row, invoice):
                        raise WorkbookError("Invoice number matches but its date or amount conflicts")
                    results.append(RecordResult(invoice.number, row, False))
                    continue
                legacy = [
                    row for row in range(3, sheet.max_row + 1)
                    if sheet.cell(row, 6).value is None and _matches_row(sheet, row, invoice)
                ]
                if invoice.legacy_row is not None:
                    if invoice.legacy_row not in legacy:
                        raise WorkbookError("Confirmed historical row no longer matches the invoice")
                    row = invoice.legacy_row
                    sheet.cell(row, 6, invoice.number).data_type = "s"
                    results.append(RecordResult(invoice.number, row, False))
                    continue
                if legacy and not invoice.confirmed_new:
                    raise WorkbookError(
                        f"Possible historical duplicate in rows {legacy}; review before recording"
                    )
                row = max(sheet.max_row + 1, 3)
                values = (
                    invoice.company, invoice.legal_name, invoice.issued,
                    float(invoice.amount), invoice.number,
                )
                for column, value in enumerate(values, start=2):
                    cell = sheet.cell(row, column, value)
                    if row > 3:
                        cell._style = copy(sheet.cell(row - 1, column)._style)
                    if isinstance(value, str):
                        cell.data_type = "s"
                sheet.cell(row, 4).number_format = "yyyy-mm-dd"
                sheet.cell(row, 5).number_format = '#,##0.00'
                results.append(RecordResult(invoice.number, row, True))
            with tempfile.NamedTemporaryFile(
                suffix=".xlsx", prefix=".invoices-", dir=self.path.parent, delete=False
            ) as stream:
                temporary = Path(stream.name)
            workbook.save(temporary)
            verification = load_workbook(temporary)
            try:
                verified_sheet = self._sheet(verification)
                for (name, coordinate), value in original_values.items():
                    if verification[name][coordinate].value != value:
                        raise WorkbookError("Save verification failed: an existing value changed")
                for invoice, result in zip(invoices, results, strict=True):
                    if (
                        not _matches_row(verified_sheet, result.row, invoice)
                        or _normalized(verified_sheet.cell(result.row, 6).value)
                        != _normalized(invoice.number)
                    ):
                        raise WorkbookError("Save verification failed for an approved invoice")
            finally:
                verification.close()
            if hashlib.sha256(self.path.read_bytes()).hexdigest() != original_hash:
                raise WorkbookError("Workbook changed during review; nothing was overwritten")
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            backup = self.backup_dir / (
                f"{self.path.stem}-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
            )
            shutil.copy2(self.path, backup)
            if hashlib.sha256(backup.read_bytes()).hexdigest() != original_hash:
                raise WorkbookError("Backup verification failed; workbook was not overwritten")
            os.replace(temporary, self.path)
            temporary = None
            return results
        finally:
            workbook.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)
