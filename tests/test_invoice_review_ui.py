import subprocess
import sys
import textwrap
import tkinter as tk

import pytest
from openpyxl import Workbook, load_workbook


@pytest.mark.parametrize("iteration", range(12))
def test_workbook_then_tk_in_same_process(tmp_path, iteration):
    path = tmp_path / f"synthetic-{iteration}.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    loaded = load_workbook(path)
    loaded.close()
    root = tk.Tk()
    try:
        root.withdraw()
        root.update_idletasks()
    finally:
        root.destroy()


def test_real_review_window_in_isolated_process(tmp_path):
    script = textwrap.dedent(r"""
        from datetime import date
        from decimal import Decimal
        from pathlib import Path
        import sys
        from tkinter import messagebox, ttk
        from openpyxl import Workbook, load_workbook
        from mail_invoices.documents import InvoiceCandidate
        from mail_invoices.review import ReviewItem
        from mail_invoices.review_ui import InvoiceReviewWindow
        from mail_invoices.settings import MailSettings

        folder = Path(sys.argv[1])
        path = folder / "synthetic.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Hoja1"
        for col, header in enumerate(("Empresa", "Raz\u00f3n Social", "Fecha", "Valor"), 2):
            sheet.cell(2, col, header)
        book.save(path)
        book.close()
        original = path.read_bytes()
        settings = MailSettings(
            folder / "opera.exe", folder / "profile", path, folder / "backups",
            "Hoja1", 1, approver="Test",
        )
        item = ReviewItem(InvoiceCandidate(
            source="synthetic.xml", digest="test", number="TEST-1",
            legal_name="Synthetic SAS", issued=date(2026, 10, 31),
            amount=Decimal("100"), currency="COP", method="xml",
        ), "Synthetic message")
        window = InvoiceReviewWindow([item], [], settings)
        window.root.withdraw()
        window.root.update_idletasks()
        assert window.forms[0]["variables"]["company"].get() == "Synthetic SAS"
        assert window.approver.get() == "Test"
        assert not window.forms[0]["selected"].get()

        def descendants(widget):
            for child in widget.winfo_children():
                yield child
                yield from descendants(child)

        widgets = list(descendants(window.root))
        buttons = {
            widget.cget("text"): widget for widget in widgets if isinstance(widget, ttk.Button)
        }
        assert any(
            str(path) in widget.cget("text")
            for widget in widgets if isinstance(widget, ttk.Label)
        )
        buttons["Seleccionar todas"].invoke()
        assert window.forms[0]["selected"].get()
        assert path.read_bytes() == original

        def unexpected_dialog(*args, **kwargs):
            raise AssertionError("Unexpected dialog")

        messagebox.askyesno = unexpected_dialog
        messagebox.showerror = unexpected_dialog
        messagebox.showinfo = lambda *args, **kwargs: None
        # Surface callback exceptions as test failures rather than Tk stderr only.
        window.root.report_callback_exception = lambda *args: unexpected_dialog()
        buttons["Aprobar y guardar lote"].invoke()
        assert window.saved
        saved = load_workbook(path)
        assert saved["Hoja1"]["F3"].value == "TEST-1"
        saved.close()
        backups = list(settings.backup_dir.glob("*.xlsx"))
        assert len(backups) == 1
        assert backups[0].read_bytes() == original
    """)
    result = subprocess.run(
        [sys.executable, "-c", script, str(tmp_path)],
        capture_output=True, text=True, timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
