from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from mail_invoices.workbook import InvoiceWorkbook, ReviewedInvoice, WorkbookError


@pytest.fixture
def store(tmp_path):
    path = tmp_path / "invoices.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    for column, header in enumerate(["Empresa", "Raz\u00f3n Social", "Fecha", "Valor"], 2):
        sheet.cell(2, column, header)
    sheet["A1"] = "=1+2"
    workbook.save(path)
    return InvoiceWorkbook(path, tmp_path / "backups")


def invoice(**changes):
    values = dict(
        company="Example", legal_name="Example SAS", issued=date(2024, 1, 2),
        amount=Decimal("1234.56"), number="TEST-001", currency="COP",
    )
    values.update(changes)
    return ReviewedInvoice(**values)


def test_approved_commit_adds_number_column_and_preserves_existing_cells(store):
    result = store.record([invoice()])
    assert result[0].inserted
    workbook = load_workbook(store.path)
    sheet = workbook["Hoja1"]
    assert sheet["F2"].value == "N\u00famero de factura"
    assert sheet["F3"].value == "TEST-001"
    assert sheet["A1"].value == "=1+2"
    assert sheet["D3"].value == datetime(2024, 1, 2)
    assert len(list(store.backup_dir.glob("*.xlsx"))) == 1
    workbook.close()


def test_number_plus_supplier_prevents_duplicate_rows(store):
    store.record([invoice()])
    result = store.record([invoice()])
    assert not result[0].inserted
    assert result[0].row == 3
    assert store.record([invoice(legal_name="Another SAS")])[0].row == 4


def test_conflicting_duplicate_does_not_modify_the_workbook(store):
    store.record([invoice()])
    original = store.path.read_bytes()
    with pytest.raises(WorkbookError, match="conflicts"):
        store.record([invoice(amount=Decimal("12"))])
    assert store.path.read_bytes() == original


def test_legacy_match_requires_explicit_review(store):
    workbook = load_workbook(store.path)
    sheet = workbook["Hoja1"]
    sheet.append([None, "Example", "Example SAS", date(2024, 1, 2), 1234.56])
    workbook.save(store.path)
    workbook.close()
    with pytest.raises(WorkbookError, match="historical duplicate"):
        store.record([invoice()])
    result = store.record([invoice(legacy_row=3)])
    assert result[0].row == 3 and not result[0].inserted


def test_spreadsheet_formula_injection_is_stored_as_text(store):
    store.record([invoice(company="=HYPERLINK(\"https://invalid\")", number="=123")])
    workbook = load_workbook(store.path)
    assert workbook["Hoja1"]["B3"].data_type == "s"
    assert workbook["Hoja1"]["F3"].data_type == "s"
    workbook.close()


def test_other_currencies_are_not_silently_mixed_in_valor(store):
    with pytest.raises(WorkbookError, match="COP"):
        store.record([invoice(currency="USD")])


@pytest.mark.parametrize("number_format", [
    '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@_-',
    '"$" #,##0', "#,##0.000", "General",
])
def test_new_amounts_keep_previous_numeric_row_format_across_batch(store, number_format):
    workbook = load_workbook(store.path)
    sheet = workbook["Hoja1"]
    sheet["E3"] = 42
    sheet["E3"].number_format = number_format
    sheet["B4"] = "Blank amount"
    sheet["E4"].number_format = "0%"
    workbook.save(store.path)
    workbook.close()
    store.record([invoice(), invoice(number="TEST-002", amount=Decimal("55"))])
    workbook = load_workbook(store.path)
    try:
        sheet = workbook["Hoja1"]
        assert sheet["E3"].number_format == number_format
        assert sheet["E4"].number_format == "0%"
        assert sheet["E5"].number_format == number_format
        assert sheet["E6"].number_format == number_format
        assert sheet["E5"].value == 1234.56
        assert sheet["E6"].value == 55
    finally:
        workbook.close()


def test_empty_workbook_uses_default_numeric_format(store):
    store.record([invoice()])
    workbook = load_workbook(store.path)
    try:
        assert workbook["Hoja1"]["E3"].number_format == "#,##0.00"
    finally:
        workbook.close()


def test_explicit_format_restoration_keeps_values_and_creates_backup(store):
    workbook = load_workbook(store.path)
    sheet = workbook["Hoja1"]
    sheet["E3"] = 42
    sheet["E3"].number_format = '"$" #,##0.00'
    workbook.save(store.path)
    workbook.close()
    result = store.record([invoice()])[0]
    workbook = load_workbook(store.path)
    sheet = workbook["Hoja1"]
    sheet.cell(result.row, 5).number_format = "#,##0.00"
    workbook.save(store.path)
    workbook.close()
    before = store.path.read_bytes()
    result = store.record([invoice()], restore_amount_format=True)[0]
    assert not result.inserted
    workbook = load_workbook(store.path)
    try:
        assert workbook["Hoja1"].max_row == 4
        assert workbook["Hoja1"]["E4"].value == 1234.56
        assert workbook["Hoja1"]["E4"].number_format == '"$" #,##0.00'
    finally:
        workbook.close()
    assert any(path.read_bytes() == before for path in store.backup_dir.glob("*.xlsx"))


def test_restoration_does_not_insert_an_unknown_invoice(store):
    before = store.path.read_bytes()
    with pytest.raises(WorkbookError, match="existing exact"):
        store.record([invoice()], restore_amount_format=True)
    assert store.path.read_bytes() == before
