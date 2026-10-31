from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from mail_invoices.documents import InvoiceCandidate
from mail_invoices.main import _prefer_xml_candidates
from mail_invoices.review_ui import (
    InvoiceReviewWindow,
    ReviewItem,
    record_approved_batch,
    reviewed_invoice_from_fields,
    suggested_company,
)
from mail_invoices.settings import MailSettings, load_settings
from mail_invoices.workbook import InvoiceWorkbook


def test_review_fields_are_converted_and_validated():
    invoice = reviewed_invoice_from_fields(review_fields())
    assert invoice.issued == date(2026, 10, 27)
    assert invoice.amount == Decimal("50900.00")
    assert invoice.currency == "COP"


def review_fields():
    return {
        "company": "Mercado",
        "legal_name": "Proveedor SAS",
        "issued": "2026-10-27",
        "amount": "50.900,00",
        "number": "FH4640622",
        "currency": "cop",
    }


def test_review_fields_reject_unsupported_currency():
    fields = review_fields()
    fields["currency"] = "USD"
    with pytest.raises(ValueError, match="Only confirmed COP"):
        reviewed_invoice_from_fields(fields)


def test_record_requires_explicit_batch_approval_and_uses_only_synthetic_workbook(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    for column, header in enumerate(["Empresa", "Razón Social", "Fecha", "Valor"], 2):
        sheet.cell(2, column, header)
    workbook.save(path)
    workbook.close()
    store = InvoiceWorkbook(path, tmp_path / "backups")
    invoice = reviewed_invoice_from_fields(review_fields())
    original = path.read_bytes()

    assert record_approved_batch(store, [invoice], approved=False) is None
    assert path.read_bytes() == original

    results = record_approved_batch(store, [invoice], approved=True)
    assert results is not None
    assert results[0].inserted
    saved = load_workbook(path)
    assert saved["Hoja1"]["F3"].value == "FH4640622"
    saved.close()


def test_zip_pdf_duplicate_prefers_xml_candidate():
    pdf = ReviewItem(
        InvoiceCandidate(
            source="tmp/invoice.zip:invoice.pdf",
            digest="pdf",
            number="FE-123",
            method="pdf_text",
        ),
        "message",
    )
    xml = ReviewItem(
        InvoiceCandidate(
            source="tmp/invoice.zip:invoice.xml",
            digest="xml",
            number="FE-123",
            method="xml",
        ),
        "message",
    )
    other = ReviewItem(
        InvoiceCandidate(
            source="tmp/other.zip:other.xml",
            digest="other",
            number="FE-456",
            method="xml",
        ),
        "message",
    )
    assert _prefer_xml_candidates([pdf, xml, other]) == [xml, other]


def test_company_suggestions_use_only_unanimous_history(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    for column, header in enumerate(["Empresa", "Raz\u00f3n Social", "Fecha", "Valor"], 2):
        sheet.cell(2, column, header)
    rows = [
        ("Mercado", "Proveedor  SAS"),
        ("Mercado", "proveedor sas"),
        ("Salud", "Clinica SAS"),
        ("Hogar", "Clinica SAS"),
        (None, "Sin Empresa SAS"),
    ]
    for row, (company, legal_name) in enumerate(rows, 3):
        sheet.cell(row, 2, company)
        sheet.cell(row, 3, legal_name)
    workbook.save(path)
    workbook.close()
    original = path.read_bytes()

    suggestions = InvoiceWorkbook(path, tmp_path / "backups").company_suggestions()

    assert suggested_company(suggestions, "PROVEEDOR SAS") == "Mercado"
    assert suggested_company(suggestions, "Clinica SAS") == "Clinica SAS"
    assert suggested_company(suggestions, "Sin Empresa SAS") == "Sin Empresa SAS"
    assert suggested_company(suggestions, "  ") == ""
    assert path.read_bytes() == original


def test_approver_is_loaded_from_config(tmp_path):
    config = tmp_path / "config.yaml"
    config.write_text(
        "opera_executable: opera.exe\n"
        "profile_dir: profile\n"
        "workbook: book.xlsx\n"
        "backup_dir: backups\n"
        "approver: '  Santiago   Arbelaez '\n",
        encoding="utf-8",
    )
    assert load_settings(config).approver == "Santiago Arbelaez"


@pytest.fixture
def review_window(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    for column, header in enumerate(["Empresa", "Razón Social", "Fecha", "Valor"], 2):
        sheet.cell(2, column, header)
    workbook.save(path)
    workbook.close()
    settings = MailSettings(
        executable=tmp_path / "opera.exe",
        profile_dir=tmp_path / "profile",
        workbook=path,
        backup_dir=tmp_path / "backups",
        sheet="Hoja1",
        batch_limit=2,
        approver="Test",
    )
    items = [
        ReviewItem(
            InvoiceCandidate(
                source=f"synthetic-{number}.xml",
                digest=number,
                number=number,
                legal_name="Synthetic SAS",
                issued=date(2026, 10, 27),
                amount=Decimal("100"),
                currency="COP",
                method="xml",
            ),
            "Synthetic message",
        )
        for number in ("TEST-1", "TEST-2")
    ]
    window = InvoiceReviewWindow(items, [], settings)
    window.root.withdraw()
    yield window
    if not window.saved:
        window.root.destroy()


def test_select_all_does_not_write_and_allows_deselecting(review_window):
    window = review_window
    original = window.settings.workbook.read_bytes()
    assert all(not form["selected"].get() for form in window.forms)
    window.select_all()
    assert all(form["selected"].get() for form in window.forms)
    window.forms[1]["selected"].set(False)
    assert not window.forms[1]["selected"].get()
    assert window.settings.workbook.read_bytes() == original
    assert not window.settings.backup_dir.exists()


def test_approval_button_writes_selected_rows_without_final_dialog(review_window, monkeypatch):
    from tkinter import messagebox

    window = review_window
    original = window.settings.workbook.read_bytes()
    window.select_all()
    window.forms[1]["selected"].set(False)

    def unexpected_dialog(*args, **kwargs):
        pytest.fail("Unexpected confirmation or validation error")

    monkeypatch.setattr(messagebox, "askyesno", unexpected_dialog)
    monkeypatch.setattr(messagebox, "showerror", unexpected_dialog)
    monkeypatch.setattr(messagebox, "showinfo", lambda *args, **kwargs: None)
    window.approve_and_save()
    assert window.saved
    saved = load_workbook(window.settings.workbook)
    try:
        assert saved["Hoja1"]["F3"].value == "TEST-1"
        assert saved["Hoja1"].max_row == 3
    finally:
        saved.close()
    backups = list(window.settings.backup_dir.glob("*.xlsx"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == original


@pytest.mark.parametrize("invalid", ["selection", "approver", "amount", "legacy"])
def test_approval_controls_still_block_unapproved_or_invalid_rows(
    review_window, monkeypatch, invalid
):
    from tkinter import messagebox

    window = review_window
    original = window.settings.workbook.read_bytes()
    window.select_all()
    if invalid == "selection":
        for form in window.forms:
            form["selected"].set(False)
    elif invalid == "approver":
        window.approver.set("")
    elif invalid == "amount":
        window.forms[0]["variables"]["amount"].set("0")
    else:
        monkeypatch.setattr(window.store, "legacy_matches", lambda invoice: (3,))
    errors = []
    confirmations = []

    def reject_confirmation(title, *args, **kwargs):
        confirmations.append(title)
        return False

    monkeypatch.setattr(messagebox, "askyesno", reject_confirmation)
    monkeypatch.setattr(messagebox, "showerror", lambda *args, **kwargs: errors.append(args))
    window.approve_and_save()
    assert not window.saved
    assert window.settings.workbook.read_bytes() == original
    assert not window.settings.backup_dir.exists()
    if invalid == "legacy":
        assert confirmations == ["Posible duplicado histórico", "Confirmar factura nueva"]
    else:
        assert errors
        assert not confirmations
