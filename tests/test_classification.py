from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from mail_invoices.classification import classification_ready
from mail_invoices.outlook_folders import folder_paths
from mail_invoices.outlook_folders import read_history
from mail_invoices.workbook import InvoiceWorkbook


@pytest.mark.parametrize(
    ("statuses", "errors", "ready"),
    [
        (["registered"], [], True),
        (["registered", "registered"], [], True),
        ([], [], False),
        (["registered", "not_registered"], [], False),
        (["historical_review"], [], False),
        (["conflict"], [], False),
        (["incomplete"], [], False),
        (["registered"], ["encrypted PDF"], False),
    ],
)
def test_only_fully_registered_messages_without_errors_are_ready(statuses, errors, ready):
    assert classification_ready(statuses, errors) is ready


def test_folder_tree_preserves_paths_and_ignores_favorites():
    rows = [
        {"id": "favoritesRoot", "level": "1", "title": "Favoritos", "name": "favoritos"},
        {"id": "", "level": "2", "title": "Exito", "name": "exito"},
        {"id": "primaryMailboxRoot_test", "level": "1", "title": "Mailbox", "name": "mailbox"},
        {"id": "", "level": "2", "title": "Bandeja de entrada : Elementos 50", "name": "bandeja de entrada"},
        {"id": "", "level": "2", "title": "Compras : Elementos 0", "name": "compras"},
        {"id": "", "level": "3", "title": "Exito : Elementos 10", "name": "exito"},
        {"id": "", "level": "2", "title": "Viajes", "name": "viajes"},
        {"id": "", "level": "3", "title": "Exito", "name": "exito"},
        {"id": "", "level": "1", "title": "", "name": "ir a grupos"},
    ]
    assert [f.path for f in folder_paths(rows)] == ["Compras", "Compras/Exito", "Viajes", "Viajes/Exito"]


def test_missing_mailbox_is_rejected():
    with pytest.raises(LookupError, match="missing"):
        folder_paths([])


def test_no_authorized_history_does_not_read_page():
    assert read_history(None, []) == []


def test_invalid_history_limit_is_rejected():
    with pytest.raises(ValueError, match="limit"):
        read_history(None, [], limit=6)


@pytest.mark.parametrize(
    ("number", "issued", "amount", "expected"),
    [
        ("TEST-1", date(2026, 10, 31), Decimal("100"), "registered"),
        ("TEST-1", date(2026, 10, 31), Decimal("200"), "conflict"),
        ("NEW", date(2026, 10, 31), Decimal("100"), "historical_review"),
        ("NEW", date(2026, 10, 30), Decimal("100"), "not_registered"),
        ("", date(2026, 10, 31), Decimal("100"), "incomplete"),
    ],
)
def test_registration_checks_are_readonly(tmp_path, number, issued, amount, expected):
    path = tmp_path / "synthetic.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Hoja1"
    for column, header in enumerate(("Empresa", "Razón Social", "Fecha", "Valor", "Número de factura"), 2):
        sheet.cell(2, column, header)
    for row, num in ((3, "TEST-1"), (4, None)):
        for col, value in enumerate(("Salud", "Synthetic SAS", date(2026, 10, 31), 100, num), 2):
            sheet.cell(row, col, value)
    book.save(path)
    book.close()
    original = path.read_bytes()
    store = InvoiceWorkbook(path, tmp_path / "backups")
    assert store.registration_status(number, "Synthetic SAS", issued, amount, "COP") == expected
    assert path.read_bytes() == original
    assert not (tmp_path / "backups").exists()
