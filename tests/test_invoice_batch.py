from dataclasses import asdict, replace
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from mail_invoices import batch_workflow
from mail_invoices.batch_state import BatchMessage, BatchStateStore, batch_lock, batch_scope
from mail_invoices.classification_decisions import ClassificationDecisionStore
from mail_invoices.documents import InvoiceCandidate
from mail_invoices.folders import IncomingMessage
from mail_invoices.review import ReviewItem
from mail_invoices.settings import MailSettings
from mail_invoices.workbook import InvoiceWorkbook, ReviewedInvoice


def pending(number="TEST-1"):
    candidate = InvoiceCandidate(
        "synthetic.xml", "digest", number, "Synthetic SAS", date(2026, 10, 31),
        Decimal("100"), "COP", method="xml", descriptions=("Consulta general",),
    )
    return BatchMessage(
        number, "sender@example.invalid", f"Factura {number}",
        [ReviewItem(candidate, f"Factura {number}")],
        IncomingMessage(
            "sender@example.invalid", f"Factura {number}", "Synthetic SAS", "Consulta",
            ({"supplier": "Synthetic SAS", "descriptions": "Consulta general", "evidence": ""},),
        ), folder="Salud",
    )


def approved(number="TEST-1"):
    return replace(
        pending(number), state="approved", approver="Santi", approved_at="2026-10-31T00:00:00+00:00",
        invoices=[ReviewedInvoice(
            "Synthetic", "Synthetic SAS", date(2026, 10, 31), Decimal("100"), number, "COP",
        )],
    )


@pytest.fixture
def stores(tmp_path):
    path = tmp_path / "book.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Hoja1"
    for column, header in enumerate(("Empresa", "Razón Social", "Fecha", "Valor"), 2):
        sheet.cell(2, column, header)
    book.save(path)
    book.close()
    return (
        BatchStateStore(tmp_path / "state.sqlite3", "test"),
        InvoiceWorkbook(path, tmp_path / "backups"),
        ClassificationDecisionStore(tmp_path / "decisions.sqlite3"),
    )


def test_state_survives_reopening_and_is_scoped(stores):
    journal, _, _ = stores
    journal.save([pending(), approved("TEST-2")])
    loaded = BatchStateStore(journal.path, "test").load()
    assert loaded == [pending(), approved("TEST-2")]
    assert BatchStateStore(journal.path, "other-mailbox").load() == []
    updated = replace(loaded[1], state="registered")
    journal.save([updated])
    assert {message.conversation: message.state for message in journal.load()} == {
        "TEST-1": "pending", "TEST-2": "registered",
    }


def test_invalid_approval_batch_is_atomic(stores):
    journal, _, _ = stores
    with pytest.raises(ValueError):
        journal.save([approved(), replace(approved("TEST-2"), approver="")])
    assert journal.load() == []
    with pytest.raises(ValueError):
        journal.save([replace(approved(), errors=["unreadable attachment"])])


def test_scope_changes_when_workbook_or_profile_changes(tmp_path):
    settings = MailSettings(tmp_path / "opera", tmp_path / "profile", tmp_path / "book", tmp_path, "Hoja1", 3)
    assert batch_scope(settings) != batch_scope(replace(settings, workbook=tmp_path / "other"))
    assert batch_scope(settings) != batch_scope(replace(settings, profile_dir=tmp_path / "other"))


def test_lock_prevents_concurrent_execution_and_releases_after_exception(tmp_path):
    path = tmp_path / "batch.lock"
    with pytest.raises(RuntimeError):
        with batch_lock(path):
            with pytest.raises(ValueError, match="ejecucion"):
                with batch_lock(path):
                    pytest.fail("Concurrent batch acquired lock")
            raise RuntimeError("Synthetic failure")
    with batch_lock(path):
        pass


def mock_folders(monkeypatch):
    monkeypatch.setattr(batch_workflow, "discover_folders", lambda page: [SimpleNamespace(path="Salud")])


def test_approved_execution_records_and_persists_before_move(stores, monkeypatch):
    journal, workbook, decisions = stores
    message = approved()
    journal.save([message])
    mock_folders(monkeypatch)
    states = []

    def move(page, data, folder, *, approved, before_move):
        assert approved and data["eligible"]
        assert journal.load()[0].state == "registered"
        assert workbook.registration_status(
            "TEST-1", "Synthetic SAS", date(2026, 10, 31), Decimal("100"), "COP",
        ) == "registered"
        before_move()
        states.append(journal.load()[0].state)

    monkeypatch.setattr(batch_workflow, "move_registered_message", move)
    batch_workflow.execute_message(message, journal, workbook, decisions, None)
    assert states == ["moving"]
    assert journal.load()[0].state == "moved"
    assert decisions.load()[0].folder == "Salud"


def test_registered_resume_does_not_write_excel_again(stores, monkeypatch):
    journal, workbook, decisions = stores
    message = approved()
    workbook.record(message.invoices)
    message.state = "registered"
    journal.save([message])
    before = workbook.path.read_bytes()
    mock_folders(monkeypatch)
    monkeypatch.setattr(workbook, "record", lambda invoices: pytest.fail("Excel rewritten"))
    monkeypatch.setattr(
        batch_workflow, "move_registered_message",
        lambda page, data, folder, *, approved, before_move: before_move(),
    )
    batch_workflow.execute_message(journal.load()[0], journal, workbook, decisions, None)
    assert journal.load()[0].state == "moved"
    assert workbook.path.read_bytes() == before


def test_crash_after_excel_before_state_save_does_not_duplicate(stores, monkeypatch):
    journal, workbook, decisions = stores
    message = approved()
    journal.save([message])
    # Represents a crash after record() committed but before registered was persisted.
    workbook.record(message.invoices)
    mock_folders(monkeypatch)
    monkeypatch.setattr(
        batch_workflow, "move_registered_message",
        lambda page, data, folder, *, approved, before_move: before_move(),
    )
    batch_workflow.execute_message(journal.load()[0], journal, workbook, decisions, None)
    book = load_workbook(workbook.path)
    try:
        assert book["Hoja1"].max_row == 3
    finally:
        book.close()
    assert journal.load()[0].state == "moved"


@pytest.mark.parametrize("verified", [True, False])
def test_uncertain_movement_resume_only_verifies_never_repeats(stores, monkeypatch, verified):
    journal, workbook, decisions = stores
    message = replace(approved(), state="moving")
    workbook.record(message.invoices)
    journal.save([message])
    mock_folders(monkeypatch)
    monkeypatch.setattr(batch_workflow, "move_registered_message", lambda *a, **kw: pytest.fail("Move repeated"))

    def verify(page, data, folder):
        if not verified:
            raise LookupError("Destination not verified")

    monkeypatch.setattr(batch_workflow, "verify_destination", verify)
    if verified:
        batch_workflow.execute_message(journal.load()[0], journal, workbook, decisions, None)
        assert journal.load()[0].state == "moved"
    else:
        with pytest.raises(LookupError):
            batch_workflow.execute_message(journal.load()[0], journal, workbook, decisions, None)
        assert journal.load()[0].state == "moving"
        assert journal.load()[0].last_error == "Destination not verified"


def test_failed_move_after_click_is_durable_and_not_retried(stores, monkeypatch):
    journal, workbook, decisions = stores
    message = approved()
    journal.save([message])
    mock_folders(monkeypatch)

    def move(page, data, folder, *, approved, before_move):
        before_move()
        raise LookupError("Uncertain result")

    monkeypatch.setattr(batch_workflow, "move_registered_message", move)
    with pytest.raises(LookupError):
        batch_workflow.execute_message(message, journal, workbook, decisions, None)
    assert journal.load()[0].state == "moving"


def test_registration_conflict_blocks_move(stores, monkeypatch):
    journal, workbook, decisions = stores
    message = replace(approved(), state="registered")
    journal.save([message])
    monkeypatch.setattr(batch_workflow, "move_registered_message", lambda *a, **kw: pytest.fail("Move attempted"))
    with pytest.raises(ValueError, match="registro"):
        batch_workflow.execute_message(message, journal, workbook, decisions, None)
    assert journal.load()[0].state == "blocked"
    assert not decisions.path.exists()


@pytest.mark.parametrize("state", ["pending", "blocked", "moved"])
def test_non_executable_states_have_no_side_effects(stores, state):
    journal, _, _ = stores
    message = replace(approved(), state=state)
    batch_workflow.execute_message(message, journal, None, None, None)
    assert not journal.path.exists()


def test_preparation_does_not_store_raw_evidence_or_paths():
    message = pending()
    candidate = replace(message.items[0].candidate, source=r"C:\private\invoice.xml", evidence="Sensitive raw text")
    raw = {
        "conversation": "TEST", "sender": message.sender, "subject": message.subject,
        "errors": [], "review_items": [ReviewItem(candidate, message.subject)],
        "classification_context": asdict(message.context),
    }
    prepared = BatchMessage.from_report(raw)
    assert prepared.items[0].candidate.evidence == ""
    assert prepared.items[0].candidate.source == "invoice.xml"


def test_blocked_records_do_not_starve_ready_pending_messages():
    blocked = replace(pending("BLOCKED"), state="blocked")
    error = replace(pending("ERROR"), errors=["Unreadable attachment"])
    ready = pending("READY")
    done = replace(approved("DONE"), state="moved")
    assert batch_workflow._active_messages([blocked, error, done, ready], 1) == [ready]


@pytest.mark.parametrize("resume", [True, False])
def test_batch_orchestrator_limits_execution_and_resumes_without_scan(
    stores, monkeypatch, tmp_path, resume,
):
    journal, workbook, _ = stores
    settings = MailSettings(tmp_path / "opera", tmp_path, workbook.path, tmp_path / "backups", "Hoja1", 20)
    journal = BatchStateStore(journal.path, batch_scope(settings))
    journal.save([approved(f"TEST-{index}") for index in range(1, 5)])
    monkeypatch.setattr(batch_workflow, "batch_state_path", lambda: journal.path)
    monkeypatch.setattr(batch_workflow, "decisions_path", lambda: tmp_path / "decisions.sqlite3")
    scanned = []

    def simulate(config, days, history, use_model, *, include_review_items):
        assert config.batch_limit == 2 and include_review_items
        scanned.append(True)
        return {"folders": ["Salud"], "messages": []}

    monkeypatch.setattr(batch_workflow, "simulate_classification", simulate)
    mock_folders(monkeypatch)
    closed = []
    page = SimpleNamespace(wait_for_selector=lambda *a, **kw: None)
    context = SimpleNamespace(close=lambda: closed.append(True))

    class Session:
        def __enter__(self):
            return SimpleNamespace(chromium=SimpleNamespace(
                launch_persistent_context=lambda *a, **kw: context,
            ))

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(batch_workflow, "sync_playwright", Session)
    monkeypatch.setattr(batch_workflow, "outlook_page", lambda ctx: page)

    class Review:
        def __init__(self, messages, folders, config, journal):
            assert len(messages) == 2

        def run(self):
            return True

    monkeypatch.setattr(batch_workflow, "BatchReviewWindow", Review)
    executed = []

    def execute(message, journal, store, decisions, page):
        executed.append(message.conversation)
        message.state = "moved"
        journal.save([message])

    monkeypatch.setattr(batch_workflow, "execute_message", execute)
    assert batch_workflow.process_invoice_batch(settings, 30, False, 2, resume=resume) == 0
    assert executed == ["TEST-1", "TEST-2"]
    assert len(scanned) == (0 if resume else 1)
    assert closed == [True]
    assert len([m for m in journal.load() if m.state == "approved"]) == 2


@pytest.mark.parametrize("cancel", [True, False])
def test_batch_cancellation_or_failure_stops_other_durable_approvals(
    stores, monkeypatch, tmp_path, cancel,
):
    journal, workbook, _ = stores
    settings = MailSettings(tmp_path / "opera", tmp_path, workbook.path, tmp_path, "Hoja1", 3)
    journal = BatchStateStore(journal.path, batch_scope(settings))
    journal.save([approved("TEST-1"), approved("TEST-2")])
    original = workbook.path.read_bytes()
    monkeypatch.setattr(batch_workflow, "batch_state_path", lambda: journal.path)
    monkeypatch.setattr(batch_workflow, "decisions_path", lambda: tmp_path / "decisions.sqlite3")
    mock_folders(monkeypatch)
    closed, attempted = [], []
    page = SimpleNamespace(wait_for_selector=lambda *a, **kw: None)
    context = SimpleNamespace(close=lambda: closed.append(True))

    class Session:
        def __enter__(self):
            return SimpleNamespace(chromium=SimpleNamespace(
                launch_persistent_context=lambda *a, **kw: context,
            ))

        def __exit__(self, *args):
            return False

    class Review:
        def __init__(self, *args):
            pass

        def run(self):
            return not cancel

    monkeypatch.setattr(batch_workflow, "sync_playwright", Session)
    monkeypatch.setattr(batch_workflow, "outlook_page", lambda ctx: page)
    monkeypatch.setattr(batch_workflow, "BatchReviewWindow", Review)

    def execute(message, journal, store, decisions, page):
        attempted.append(message.conversation)
        message.last_error = "Synthetic failure"
        journal.save([message])
        raise LookupError("Synthetic failure")

    monkeypatch.setattr(batch_workflow, "execute_message", execute)
    assert batch_workflow.process_invoice_batch(settings, 30, False, 3, resume=True) == 1
    assert attempted == ([] if cancel else ["TEST-1"])
    assert journal.load()[1].state == "approved"
    assert workbook.path.read_bytes() == original
    assert closed == [True]


@pytest.mark.parametrize("flag", ["--process-invoice-batch", "--resume-invoice-batch"])
def test_cli_routes_batch_modes(monkeypatch, tmp_path, flag):
    from mail_invoices import main

    settings = MailSettings(tmp_path / "opera", tmp_path, tmp_path / "book", tmp_path, "Hoja1", 20)
    monkeypatch.setattr(main, "load_settings", lambda path: settings)
    monkeypatch.setattr("sys.argv", ["mail-invoices", flag, "--batch-size", "2", "--days", "7", "--no-model"])
    calls = []
    monkeypatch.setattr(
        batch_workflow, "process_invoice_batch",
        lambda *args, **kw: calls.append((args, kw)) or 0,
    )
    assert main.main() == 0
    assert calls == [((settings, 7, False, 2), {"resume": flag == "--resume-invoice-batch"})]


@pytest.mark.parametrize("size", [0, 6])
def test_cli_rejects_unsafe_batch_sizes(monkeypatch, size):
    from mail_invoices import main

    monkeypatch.setattr("sys.argv", ["mail-invoices", "--process-invoice-batch", "--batch-size", str(size)])
    with pytest.raises(SystemExit) as error:
        main.main()
    assert error.value.code == 2
